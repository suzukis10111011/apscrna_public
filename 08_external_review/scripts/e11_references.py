#!/usr/bin/env python
"""Purpose: Take up the rows the second external check marked "numerical agreement only,
  reference mismatch": the printed value was supported, but the cell the worksheet named
  held another quantity. Rows whose quantity does exist in one cell are pointed at it and
  read back; rows where no single cell can carry the quantity are demoted to a sentence
  check with the reason and, where one exists, the derivation.
Input: 07_submission_bundle/tables/Supplementary_Tables_S1-S12.xlsx,
  07_submission_bundle/tables/manuscript_number_agreement.csv,
  07_submission_bundle/tables/check_report_v8_ChatGPT_20261007.txt (the remarks)
Output: 08_external_review/tables/E11_reference_corrections.csv

Targets are given as identifying column/value pairs and the column wanted, never as a
literal address, so the table stays correct when a part is added above the row. Each
resolved cell is read back against the value the manuscript prints, at the precision it
prints it; a target that does not read back is reported rather than written.
"""
import decimal
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

W = os.environ.get("APSC_ROOT", "/path/to/Documents/apscrna_v2")
BUN = f"{W}/07_submission_bundle/tables"
OUT = f"{W}/08_external_review/tables"

WB = load_workbook(f"{BUN}/Supplementary_Tables_S1-S12.xlsx", data_only=True)
A = pd.read_csv(f"{BUN}/manuscript_number_agreement.csv")
A = A[A.version == sorted(v for v in A.version.unique()
                          if v.startswith("clean_"))[-1]].reset_index(drop=True)


def rnd(v, dp):
    return float(decimal.Decimal(repr(float(v))).quantize(
        decimal.Decimal(1).scaleb(-dp), rounding=decimal.ROUND_HALF_UP))


def locate(sheet, keys, column, pick=None):
    ws = WB[sheet]
    header = [c.value for c in ws[1]]
    recs = [dict(zip(header, row), _r=i)
            for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3)]
    df = pd.DataFrame(recs)
    m = pd.Series(True, index=df.index)
    for k, v in keys.items():
        col = df[k].astype(str)
        m &= (col != str(v[1])) if isinstance(v, tuple) else (col == str(v))
    sub = df[m]
    assert len(sub) >= 1, (sheet, keys, "no row")
    if pick in ("min", "max"):
        num = pd.to_numeric(sub[column], errors="coerce")
        sub = sub.loc[[num.idxmin() if pick == "min" else num.idxmax()]]
    assert len(sub) == 1, (sheet, keys, f"{len(sub)} rows")
    r = sub.iloc[0]
    return f"{get_column_letter(header.index(column) + 1)}{int(r._r)}", r[column]


DETECT = dict(table_part="detection fraction", cohort="discovery_n5")
COMPO = dict(table_part="compartment composition per lesion, both annotation versions",
             annotation_version="03c", cohort="discovery_n5", level="compartment")
DISP = dict(table_part="dispersion baselines", conditioning="RELA_activity",
            score_conditioned="lf_D_tissueforming")
DONOR = dict(table_part="per donor")

# No -> (sheet, keys, column, pick, quantity)
RETARGET = {
    28: ("S12", dict(table_part="single-donor check", lesion="CAP1"), "declared_donor",
         None, "ages of the three validation donors as the source paper declares them"),
    110: ("S1", dict(table_part="gene-set and marker-panel coverage",
                     gene_set="axis A_cytokine_chemokine"), "n_def", None,
          "size of the cytokine/chemokine gene set"),
    111: ("S1", dict(table_part="gene-set and marker-panel coverage",
                     gene_set="axis B_survival_antiapoptotic"), "n_def", None,
          "size of the survival/anti-apoptotic gene set"),
    146: ("S5", dict(table_part="projected state labels vs the within-cohort definition",
                     lesion="CAP1", labelling="ingest_projected",
                     scheme_adopted="yes"), "pct_inflammatory",
          None, "inflammatory-state percentage after the symbol correction"),
    220: ("S8", COMPO | dict(label="immune"), "pct_of_all_cells", "min",
          "lowest immune percentage over the discovery lesions, adopted annotation"),
    328: ("S11", DISP, "variance_retained", "max",
          "highest retained variance under conditioning on RELA activity"),
    371: ("S7", DETECT | dict(gene="IL1R1", state="tissue-forming"),
          "detection_fraction", "max",
          "highest IL1R1 detection fraction in tissue-forming cells"),
    374: ("S7", DETECT | dict(gene="TNFRSF1A", state="tissue-forming"),
          "detection_fraction", "min",
          "lowest TNFRSF1A detection fraction in tissue-forming cells"),
    376: ("S7", DETECT | dict(gene="TNFRSF1A", state="inflammatory"),
          "detection_fraction", "min",
          "lowest TNFRSF1A detection fraction in inflammatory cells"),
    377: ("S7", DETECT | dict(gene="TNFRSF1A", state="inflammatory"),
          "detection_fraction", "max",
          "highest TNFRSF1A detection fraction in inflammatory cells"),
    409: ("S9", DONOR | dict(group="Periodontitis", donor="POOLED"),
          "depth_controlled_rho", None, "pooled periodontitis depth-controlled rho"),
    411: ("S9", DONOR | dict(donor=("!=", "POOLED")), "depth_controlled_rho", "max",
          "highest per-donor depth-controlled rho in the gingival cohort"),
}

# No -> (why no single cell carries it, how to check it instead), English then Japanese
CLEAR_JA = {
    19: ("ソフトウェアの版で、測定値ではない", ""),
    59: ("探索コホートのランの塩基組成で、ケミストリー表は検証3ランのみを扱う", ""),
    85: ("QC の2段階の差を病変ごとに取った範囲", "表S1 の QC part で病変ごとに 1_qc_pass − 2_singlet"),
    144: ("投影された細胞が埋め込み上のどこに来るかという主張で、細胞数ではない",
          "細胞数そのものは表S5 の間質 part にある"),
    164: ("ソフトウェアの版で、測定値ではない", ""),
    179: ("2列の差を病変で最大化した値", "表S6 で abs(matched_sd − total_sd) の病変最大"),
    210: ("ソフトウェアの版で、測定値ではない", ""),
    213: ("ソフトウェアの版で、測定値ではない", ""),
    214: ("ソフトウェアの版で、測定値ではない", ""),
    237: ("科学記数法で書いた p 値の仮数部", ""),
    242: ("図番号", ""),
    260: ("図番号", ""),
    306: ("図番号", ""),
    359: ("受容体ファミリー名 IL-1 の一部で、測定値ではない", ""),
    357: ("図番号", ""),
    358: ("図番号", ""),
    362: ("図番号", ""),
    369: ("受容体ファミリー名 IL-6 の一部で、測定値ではない", ""),
    393: ("図番号", ""),
    417: ("図番号", ""),
    467: ("パネルに引いた有意水準の線で、図の設定値", ""),
    485: ("遺伝子集合×スコア法の組み合わせ数、すなわち行数", "表S4 のプール行を数える"),
}
QUANT_JA = {
    28: "検証3ドナーの年齢（原典の記載）", 110: "cytokine/chemokine 遺伝子集合のサイズ",
    111: "survival/anti-apoptotic 遺伝子集合のサイズ",
    146: "シンボル修正後の inflammatory 状態の割合",
    220: "採用注釈版における探索病変の免疫割合の最小値",
    328: "RELA 活性で統制したときの残存分散の最大値",
    371: "tissue-forming 細胞での IL1R1 検出率の最大値",
    374: "tissue-forming 細胞での TNFRSF1A 検出率の最小値",
    376: "inflammatory 細胞での TNFRSF1A 検出率の最小値",
    377: "inflammatory 細胞での TNFRSF1A 検出率の最大値",
    409: "歯周炎プールの深度統制 rho", 411: "歯肉コホートのドナー別深度統制 rho の最大値",
}
CLEAR = {
    19: ("a software version, not a measured value", ""),
    59: ("base composition of a discovery run; the chemistry table covers the three "
         "validation runs only", ""),
    85: ("a range over the lesions of a difference between two QC stages", 
         "per lesion, 1_qc_pass minus 2_singlet in the QC part of S1"),
    144: ("the claim is where the projected cells sit in the embedding, not a count",
          "the cell count itself is in the stromal part of S5"),
    164: ("a software version, not a measured value", ""),
    179: ("a maximum over the lesions of a difference between two columns",
          "max of abs(matched_sd - total_sd) over the lesions in S6"),
    210: ("a software version, not a measured value", ""),
    213: ("a software version, not a measured value", ""),
    214: ("a software version, not a measured value", ""),
    237: ("the mantissa of a p value written in scientific notation", ""),
    242: ("a figure number", ""),
    260: ("a figure number", ""),
    306: ("a figure number", ""),
    359: ("part of the receptor-family name IL-1, not a measured value", ""),
    357: ("a figure number", ""),
    358: ("a figure number", ""),
    362: ("a figure number", ""),
    369: ("part of the receptor-family name IL-6, not a measured value", ""),
    393: ("a figure number", ""),
    417: ("a figure number", ""),
    467: ("the significance threshold drawn in the panel, a figure setting", ""),
    485: ("the number of gene-set by scoring-method combinations, i.e. a count of rows",
          "count the pooled rows of S4"),
}

rows = []
for no, (sheet, keys, column, pick, why) in sorted(RETARGET.items()):
    cell, value = locate(sheet, keys, column, pick)
    printed = str(A.iloc[no - 1].printed)
    t = (printed.replace("−", "-").replace("–", "-").replace(",", "")
         .replace(" ", "").rstrip("%").split("/")[0])
    try:
        pv, base = abs(float(t)), abs(float(value))
        dp = len(t.split(".")[1]) if "." in t else 0
        ok = any(rnd(c, dp) == rnd(pv, dp) for c in (base, base * 100, base / 100))
    except (ValueError, TypeError):
        # a text cell: the printed value has to appear in it
        ok = t in str(value).replace(",", "")
    rows.append(dict(No=no, printed=printed, action=("applied: single cell, reads back"
                                                     if ok else "not applied: the cell "
                                                     "does not carry the printed value"),
                     new_sheet=sheet, new_cell=cell, new_column=column, new_value=value,
                     reads_back=ok, quantity=why, quantity_ja=QUANT_JA[no],
                     check_instead="", check_instead_ja=""))
for no, (why, how) in sorted(CLEAR.items()):
    rows.append(dict(No=no, printed=str(A.iloc[no - 1].printed),
                     action="cleared: no single cell carries this quantity",
                     new_sheet="", new_cell="", new_column="", new_value="",
                     reads_back="", quantity=why, quantity_ja=CLEAR_JA[no][0],
                     check_instead=how, check_instead_ja=CLEAR_JA[no][1]))
E11 = pd.DataFrame(rows).sort_values("No")
E11.to_csv(f"{OUT}/E11_reference_corrections.csv", index=False)
print(E11[E11.action.str.startswith(("applied", "not applied"))][
    ["No", "printed", "new_sheet", "new_cell", "new_column", "new_value", "reads_back"]
].to_string(index=False))
print("\napplied:", int(E11.action.str.startswith("applied").sum()),
      "| cleared:", int(E11.action.str.startswith("cleared").sum()),
      "| not applied:", int(E11.action.str.startswith("not applied").sum()))
