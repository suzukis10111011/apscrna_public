#!/usr/bin/env python
"""Purpose: Point the worksheet rows of the enrichment series, of the programme-score
  dip test, and of the seven rows the second external check found mis-referenced, at the
  cells that carry those quantities.
Input: 07_submission_bundle/tables/Supplementary_Tables_S1-S12.xlsx (rebuilt with the
  new parts), 07_submission_bundle/tables/manuscript_number_agreement.csv
Output: 08_external_review/tables/E9_reference_corrections.csv

Addresses are not written by hand: each target is given as a sheet, a set of key
column/value pairs that identify the row, and the column wanted. The row is located in
the rebuilt workbook and the address derived, so the table stays correct when a part is
added above it. Every resolved cell is read back and compared with the value the
manuscript prints at the precision it prints it.
"""
import decimal
import os
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

W = os.environ.get("APSC_ROOT", "/path/to/Documents/apscrna_v2")
BUN = f"{W}/07_submission_bundle/tables"
OUT = f"{W}/08_external_review/tables"
os.makedirs(OUT, exist_ok=True)

XLSX = f"{BUN}/Supplementary_Tables_S1-S12.xlsx"
WB = load_workbook(XLSX, data_only=True)
A = pd.read_csv(f"{BUN}/manuscript_number_agreement.csv")
A = A[A.version == sorted(v for v in A.version.unique()
                          if v.startswith("clean_"))[-1]].reset_index(drop=True)


def rnd(v, dp):
    return float(decimal.Decimal(repr(float(v))).quantize(
        decimal.Decimal(1).scaleb(-dp), rounding=decimal.ROUND_HALF_UP))


def sheet_frame(sheet):
    """the sheet as a frame, with the Excel row number of each record"""
    ws = WB[sheet]
    header = [c.value for c in ws[1]]
    recs = []
    for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        d = dict(zip(header, row))
        d["_excel_row"] = i
        recs.append(d)
    return pd.DataFrame(recs), header


def locate(sheet, keys, column, pick=None):
    """address of `column` in the single row of `sheet` matching `keys`"""
    df, header = sheet_frame(sheet)
    m = pd.Series(True, index=df.index)
    for k, v in keys.items():
        if k == "programme_score_is_per_lesion":
            m &= df["lesion"].astype(str) != "POOLED"
            continue
        m &= df[k].astype(str) == str(v)
    sub = df[m]
    assert len(sub) >= 1, (sheet, keys, column, "no row")
    if pick == "min":
        sub = sub.loc[[pd.to_numeric(sub[column]).idxmin()]]
    assert len(sub) == 1, (sheet, keys, column, f"{len(sub)} rows")
    r = sub.iloc[0]
    letter = get_column_letter(header.index(column) + 1)
    return f"{letter}{int(r._excel_row)}", r[column]


PRIM = dict(table_part="primary enrichment per lesion (as in Figure 2)")
DIP = dict(table_part="unimodality of the programme scores per lesion (Hartigan dip "
                      "test, as in Figure 4E)")
# No -> (sheet, row keys, column, pick, why)
TARGETS = {
    # Abstract and Results, the enrichment series of the main analysis
    3: ("S5", {**PRIM, "lesion": "POOLED_discovery"}, "stromal_pct_overall_primary",
        None, "stromal share of all cells, pooled discovery"),
    4: ("S5", {**PRIM, "lesion": "POOLED_discovery"}, "stromal_pct_top_decile_primary",
        None, "stromal share of the top decile, pooled discovery"),
    5: ("S5", {**PRIM, "lesion": "A"}, "enrichment_fold_primary", None,
        "lowest per-lesion enrichment of the eight lesions"),
    6: ("S5", {**PRIM, "lesion": "CAP2"}, "enrichment_fold_primary", None,
        "highest per-lesion enrichment of the eight lesions"),
    233: ("S5", {**PRIM, "lesion": "POOLED_discovery"}, "stromal_pct_overall_primary",
          None, "stromal share of all cells, pooled discovery"),
    234: ("S5", {**PRIM, "lesion": "POOLED_discovery"}, "enrichment_fold_primary", None,
          "pooled discovery enrichment"),
    235: ("S5", {**PRIM, "lesion": "A"}, "enrichment_fold_primary", None,
          "lowest discovery per-lesion enrichment"),
    236: ("S5", {**PRIM, "lesion": "E"}, "enrichment_fold_primary", None,
          "highest discovery per-lesion enrichment"),
    240: ("S5", {**PRIM, "lesion": "CAP1"}, "enrichment_fold_primary", None,
          "lowest validation per-lesion enrichment"),
    241: ("S5", {**PRIM, "lesion": "CAP2"}, "enrichment_fold_primary", None,
          "highest validation per-lesion enrichment"),
    # programme-score unimodality, Figure 4E
    # the bound the sentence states is the smallest dip p over the 16 per-lesion tests;
    # the count of tests (No.465) is the number of rows of the part, not a cell value
    356: ("S3", {**DIP, "programme_score_is_per_lesion": True},
          "dip_p_programme_score", "min",
          "smallest dip p over the 16 per-lesion tests"),
    466: ("S3", {**DIP, "programme_score_is_per_lesion": True},
          "dip_p_programme_score", "min",
          "smallest dip p over the 16 per-lesion tests, as quoted in the caption"),
    # the seven rows the second check found pointing at another quantity
    197: ("S4", {"lesion": "POOLED"}, "depth_controlled_rho", "min",
          "minimum depth-controlled pooled rho over the gene-set and scoring variants"),
    304: ("S2", {"axis": "lf_D_tissueforming", "TF": "REL"}, "pooled_rho", None,
          "REL against the tissue-forming programme"),
    313: ("S3", {"table_part": "main", "lesion": "CAP1", "cell_set": "stromal"},
          "depth_controlled_rho_composite", None, "composite depth-controlled rho, CAP1"),
    314: ("S3", {"table_part": "main", "lesion": "CAP2", "cell_set": "stromal"},
          "depth_controlled_rho_composite", None, "composite depth-controlled rho, CAP2"),
    315: ("S3", {"table_part": "main", "lesion": "CAP3", "cell_set": "stromal"},
          "depth_controlled_rho_composite", None, "composite depth-controlled rho, CAP3"),
    384: ("S10", {"table_part": "per-lesion and per-state statistics",
                  "cohort": "validation", "inflammatory_set": "infl_new_NFkB",
                  "tissueforming_set": "tiss_new4_excl_SMAD3", "group": "lesion",
                  "label": "POOLED"}, "spearman_rho", None,
          "pooled validation TF-set correlation with SMAD3 excluded"),
    386: ("S10", {"table_part": "MYC activity against the two TF sets",
                  "cohort": "discovery_n5", "lesion": "POOLED",
                  "set_name": "tissueforming_osteo_mesenchymal",
                  "summary": "median of per-factor rho"},
          "spearman_rho_MYC_vs_set_mean", None,
          "median of the per-factor MYC correlations, tissue-forming set"),
}

rows = []
for no, (sheet, keys, column, pick, why) in sorted(TARGETS.items()):
    cell, value = locate(sheet, keys, column, pick)
    printed = str(A.iloc[no - 1].printed)
    t = (printed.replace("−", "-").replace("–", "-").replace(",", "")
         .replace(" ", "").rstrip("%").split("/")[0])
    try:
        pv = abs(float(t))
        dp = len(t.split(".")[1]) if "." in t else 0
        base = abs(float(value))
        reads_back = any(rnd(c, dp) == rnd(pv, dp) for c in (base, base * 100,
                                                             base / 100))
    except (ValueError, TypeError):
        reads_back = None
    rows.append(dict(No=no, printed=printed, new_sheet=sheet, new_cell=cell,
                     new_column=column, new_value=value, reads_back=reads_back,
                     quantity=why,
                     action=("applied: single cell, reads back" if reads_back else
                             "not applied: the cell does not carry the printed value")))
E9 = pd.DataFrame(rows)
E9.to_csv(f"{OUT}/E9_reference_corrections.csv", index=False)
print(E9[["No", "printed", "new_sheet", "new_cell", "new_column", "new_value",
          "reads_back"]].to_string(index=False))
print("\napplied:", int(E9.action.str.startswith("applied").sum()), "of", len(E9))
