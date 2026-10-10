#!/usr/bin/env python
"""Purpose: Verify PPTX slides retain editable text by checking embedded images match source PNGs, text layers
  are intact, textbox counts align, and dimensions match.
Input: none (reads PPTX files and source PNG images from project structure)
Output: {W}/results/tables/39_pptx_check.csv

39: Mechanically inspect whether each figure's pptx is in a state where "text can be edited."

There are 4 points to check: (1) whether the embedded image is identical to the plain figure PNG (MD5); (2) whether text is
baked into the embedded image (comparing the number of connected components of dark color corresponding to text size between the plain figure and the labeled version);
(3) whether the textbox count matches the same figure in the combined version, (4) whether the slide dimensions and image dimensions match.
"""
import os, glob, hashlib, pathlib, io
import numpy as np, pandas as pd
from PIL import Image
from scipy import ndimage
from pptx import Presentation
from pptx.util import Mm

import os as _os
W = _os.environ.get("APSC_ROOT") or _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
FIGD = f"{W}/results/figures_ms"
BARE = f"{FIGD}/bare"
PD = f"{FIGD}/pptx"
EMM = Mm(1).emu
ORDER = ["Fig1_design_composition", "Fig2_enrichment", "Fig3_tf_screen",
         "Fig4_coupling_coexistence", "Fig5_receptors_coactivity", "Fig6_compartments",
         "Fig7_independent_cohort", "FigS1_qc", "FigS2_integration", "FigS3_genesets_scoring",
         "FigS4_matrix_vs_requantification", "FigS5_raw_vs_depth", "FigS6_leave_one_out",
         "FigS7_state_labelling", "FigS8_other_programs"]


def md5(b):
    return hashlib.md5(b).hexdigest()


def glyphs(path_or_bytes, scale=0.25):
    """暗色の連結成分のうち、文字サイズ相当（縮小後 3-40 px、縦横比 0.15-8）の個数。"""
    im = Image.open(io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes)
    im = im.convert("L")
    im = im.resize((int(im.width * scale), int(im.height * scale)), Image.LANCZOS)
    a = np.asarray(im) < 110
    lab, n = ndimage.label(a)
    if n == 0:
        return 0
    objs = ndimage.find_objects(lab)
    cnt = 0
    for sl in objs:
        h, w = sl[0].stop - sl[0].start, sl[1].stop - sl[1].start
        if 3 <= h <= 40 and 3 <= w <= 40 and 0.15 <= (w / h) <= 8:
            cnt += 1
    return cnt


comb = {}
cp = Presentation(f"{FIGD}/figures_editable.pptx")
for i, s in enumerate(cp.slides):
    comb[ORDER[i]] = sum(1 for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())

rows = []
for name in ORDER:
    short = name.split("_")[0]
    fp = f"{PD}/{short}.pptx"
    r = dict(figure=short, pptx=os.path.basename(fp), exists=os.path.exists(fp))
    if not r["exists"]:
        rows.append(r); continue
    prs = Presentation(fp)
    s = list(prs.slides)[0]
    pics = [sh for sh in s.shapes if sh.shape_type == 13]
    tbs = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    r["n_textbox"] = len(tbs)
    r["n_textbox_combined"] = comb.get(name)
    r["textbox_match"] = (len(tbs) == comb.get(name))
    r["n_picture"] = len(pics)
    r["slide_mm"] = f"{prs.slide_width / EMM:.0f}x{prs.slide_height / EMM:.0f}"
    barep = f"{BARE}/{name}.png"
    r["bare_png"] = os.path.basename(barep)
    r["md5_bare_png"] = md5(open(barep, "rb").read())[:12] if os.path.exists(barep) else ""
    if pics:
        blob = pics[0].image.blob
        r["md5_embedded"] = md5(blob)[:12]
        r["md5_match"] = (r["md5_embedded"] == r["md5_bare_png"])
        im = Image.open(io.BytesIO(blob))
        r["image_mm"] = f"{im.width / 600 * 25.4:.0f}x{im.height / 600 * 25.4:.0f}"
        r["size_match"] = (r["image_mm"] == r["slide_mm"])
        r["glyphs_embedded"] = glyphs(blob)
    txtp = f"{FIGD}/{name}.png"
    r["glyphs_textversion"] = glyphs(txtp) if os.path.exists(txtp) else np.nan
    r["glyphs_bare_png"] = glyphs(barep) if os.path.exists(barep) else np.nan
    # The number of connected components alone is not sufficient for judgment (the plain figure includes axis labels, tick marks, and heatmap cell values
    # Intentionally left as is, so the count is comparable to that of the version with characters). Therefore, the "moved characters
    # Directly check whether dark pixels remain in the raw plate image at the recorded position.
    import json as _json
    meta = f"{BARE}/textmeta/{name}.json"
    still = 0; checked = 0
    if os.path.exists(meta) and pics:
        M = _json.load(open(meta))
        bim = np.asarray(Image.open(io.BytesIO(pics[0].image.blob)).convert("L"))
        H_, W_ = bim.shape
        tim = np.asarray(Image.open(f"{FIGD}/{name}.png").convert("L")) \
            if os.path.exists(f"{FIGD}/{name}.png") else None
        # Figure with text uses a tight crop, while the plain figure is a merge of canvas and tight, so the dimensions differ between figures
        # (Note: the same fractional coordinates point to a different location otherwise (Fig3 shifts 2% horizontally). In that case, pixel comparison
        # Use the string (drawn) remaining in the plain figure to determine a match, without using it.
        same_geom = tim is not None and tim.shape == bim.shape
        drawn_set = {str(d["text"]).strip() for d in M.get("drawn", [])}
        for itm in M["items"]:
            # tick_paren is the outline of the "retained cluster name" at the recorded position, so having dark pixels there
            # Obviously (the cluster name itself). Excluded from the judgment. The legend and panel symbols are also treated separately in terms of position.
            if itm["kind"] in ("legend", "panel_letter", "tick_paren") or "x0" not in itm:
                continue
            x0, x1 = int(itm["x0"] * W_) + 1, int(itm["x1"] * W_) - 1
            y0, y1 = int(itm["y0"] * H_) + 1, int(itm["y1"] * H_) - 1
            if x1 - x0 < 4 or y1 - y0 < 4 or x0 < 0 or y0 < 0 or x1 > W_ or y1 > H_:
                continue
            checked += 1
            db = (bim[y0:y1, x0:x1] < 110).mean()
            # Compare the same crop with the labeled version. In figures where data points or bars are present at that position,
            # since dark pixels alone cannot be used for judgment, if the ratio is below 0.6, "the text is erased".
            dt = None
            if tim is not None:
                H2, W2 = tim.shape
                tx0, tx1 = int(itm["x0"] * W2) + 1, int(itm["x1"] * W2) - 1
                ty0, ty1 = int(itm["y0"] * H2) + 1, int(itm["y1"] * H2) - 1
                if tx1 - tx0 > 4 and ty1 - ty0 > 4:
                    dt = (tim[ty0:ty1, tx0:tx1] < 110).mean()
            if same_geom and dt and dt > 0.01:
                if db / dt >= 0.6:
                    still += 1
            elif str(itm.get("text", "")).strip()[:24] in drawn_set:
                still += 1
    r["same_geometry_as_textversion"] = bool(same_geom) if pics else None
    r["moved_items_checked"] = checked
    r["moved_text_still_in_image"] = still
    r["text_burned_in"] = bool(still > 0)
    rows.append(r)

D = pd.DataFrame(rows)
D.to_csv(f'{os.environ.get("APSC_TABDIR", f"{W}/results/tables")}/39_pptx_check.csv', index=False)
cols = ["figure", "n_textbox", "n_textbox_combined", "textbox_match", "md5_match",
        "slide_mm", "image_mm", "size_match", "glyphs_embedded", "glyphs_textversion",
        "moved_items_checked", "moved_text_still_in_image", "text_burned_in"]
print(D[cols].to_string(index=False))
print("\nMD5 不一致:", D[~D.md5_match.fillna(False)].figure.tolist() or "なし")
print("テキストボックス数不一致:", D[~D.textbox_match.fillna(False)].figure.tolist() or "なし")
print("寸法不一致:", D[~D.size_match.fillna(False)].figure.tolist() or "なし")
print("画像に文字が焼き込まれている:", D[D.text_burned_in.fillna(False)].figure.tolist() or "なし")
